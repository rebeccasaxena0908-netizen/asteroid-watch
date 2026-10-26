"""FastAPI app checks that need no database: validation, errors, value conversion, exports, static files.

Live checks against MySQL are in tests/integration/test_api.py; PHP parity in tests/parity/.
"""

import csv
import hashlib
import io
import json
import re
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

import api.db
from api.db import php_value, with_bools
from api.main import SWAGGER, app
from api.params import BadRequest, like_pattern, param_date, param_int, param_text
from api.routes.approaches import approach_query
from api.routes.exports import safe_cell, to_csv, to_xlsx
from api.routes.trends import size_histogram
from api.schemas import AsteroidRow
from pipeline.config import load_settings
from pipeline.db import engine

client = TestClient(app, raise_server_exceptions=False)


@pytest.fixture
def no_database(monkeypatch):
    """Any database access fails the test: proves a request was rejected before any query."""
    def forbidden():
        raise AssertionError("the database was used")
    monkeypatch.setattr(api.db, "reader_engine", forbidden)


# Same cases as the PHP API's tests (tests/test_php.py), plus a few more.
@pytest.mark.parametrize("path, message", [
    ("/api/approaches?page=0", "'page' must be a whole number"),
    ("/api/approaches?per_page=500", "'per_page' must be a whole number from 1 to 100"),
    ("/api/approaches?sort=miss_distance_km;DROP", "'sort' must be one of"),
    ("/api/approaches?start=2026-13-01", "'start' must be a date"),
    ("/api/approaches?start=2026-10-12&end=2026-10-01", "'start' must be on or before 'end'"),
    ("/api/approaches?hazardous=yes", "'hazardous' must be one of"),
    ("/api/approaches?q=" + "x" * 51, "'q' must be at most 50 characters"),
    ("/api/asteroid?id=1%20OR%201=1", "'id' is required and must be a numeric neo_id"),
    ("/api/asteroid", "'id' is required"),
    ("/api/runs?limit=abc", "'limit' must be a whole number"),
    ("/api/trends?end=yesterday", "'end' must be a date"),
    ("/api/approaches.csv?sort=bad", "'sort' must be one of"),
])
def test_bad_input_is_rejected_before_any_query(no_database, path, message):
    r = client.get(path)
    assert r.status_code == 400
    assert message in r.json()["error"]


def test_mysql_down_gives_503_with_the_php_message(monkeypatch):
    settings = replace(load_settings(), db_socket="/nonexistent.sock", db_host="127.0.0.1", db_port=1)
    monkeypatch.setattr(api.db, "reader_engine", lambda: engine(settings, connect_timeout=1))
    r = client.get("/api/summary")
    assert r.status_code == 503
    assert r.json() == {"error": "Cannot connect to MySQL. Is XAMPP MySQL running?"}


def test_missing_read_only_user_gives_503(monkeypatch):
    monkeypatch.setattr(api.db, "load_settings", lambda: replace(load_settings(), dashboard_db_password=""))
    api.db.reader_engine.cache_clear()
    try:
        r = client.get("/api/runs")
    finally:
        api.db.reader_engine.cache_clear()
    assert r.status_code == 503
    assert "DASHBOARD_DB_PASSWORD" in r.json()["error"]


@pytest.mark.parametrize("method", ["POST", "PUT", "DELETE", "PATCH"])
def test_only_get_is_supported(no_database, method):
    r = client.request(method, "/api/summary")
    assert (r.status_code, r.json()) == (405, {"error": "Only GET is supported."})


def test_unknown_paths_are_json_404s_with_safe_headers():
    r = client.get("/api/does-not-exist")
    assert (r.status_code, r.json()) == (404, {"error": "Not found."})
    assert r.headers["Content-Type"] == "application/json; charset=utf-8"
    assert r.headers["Cache-Control"] == "no-store"
    assert r.headers["X-Content-Type-Options"] == "nosniff"


def test_pretty_output_is_indented(no_database):
    message = "'limit' must be a whole number from 1 to 50."
    assert client.get("/api/runs?limit=0&pretty=1").text == f'{{\n    "error": "{message}"\n}}'
    assert client.get("/api/runs?limit=0").text == f'{{"error":"{message}"}}'


def test_server_errors_hide_details(monkeypatch):
    def broken():
        raise RuntimeError("secret internal detail")
    monkeypatch.setattr(api.db, "reader_engine", broken)
    r = client.get("/api/summary")
    assert (r.status_code, r.json()) == (500, {"error": "Something went wrong on the server."})


def test_values_are_converted_like_php_pdo():
    assert php_value(Decimal("21.650")) == "21.650"
    assert php_value(date(2026, 10, 12)) == "2026-10-12"
    assert php_value(datetime(2026, 10, 12, 6, 30)) == "2026-10-12 06:30:00"
    assert php_value(1.5) == 1.5 and php_value(7) == 7 and php_value(None) is None
    assert with_bools({"is_potentially_hazardous": 1, "passed": 0, "is_sentry_object": None, "n": 1}) == \
        {"is_potentially_hazardous": True, "passed": False, "is_sentry_object": None, "n": 1}


def test_response_models_are_strict():
    row = {"neo_id": "1", "name": "x", "nasa_jpl_url": None, "absolute_magnitude_h": 21.65,
           "est_diameter_min_m": 1.0, "est_diameter_max_m": 2.0, "is_potentially_hazardous": False,
           "is_sentry_object": False, "first_seen_at": "2026-10-12 00:00:00", "updated_at": "2026-10-12 00:00:00"}
    with pytest.raises(ValidationError):
        AsteroidRow.model_validate(row)   # PHP sent DECIMAL as a string; a float must not slip through
    AsteroidRow.model_validate(row | {"absolute_magnitude_h": "21.650"})


@pytest.mark.parametrize("value, ok", [
    ("2026-10-12", True), ("2026-02-30", False), ("2026-1-5", False), ("20261012", False), ("yesterday", False)])
def test_param_date(value, ok):
    if ok:
        assert param_date("start", value) == value
    else:
        with pytest.raises(BadRequest):
            param_date("start", value)


def test_param_int_and_text_follow_php():
    assert param_int("page", "007", 1, 1, 10) == 7
    assert param_int("page", "", 3, 1, 10) == 3
    for bad in ["+5", "5.0", "1e2", " 5", "0", "11"]:
        with pytest.raises(BadRequest):
            param_int("page", bad, 1, 1, 10)
    assert param_text("q", " 2008 TC4\t", 50) == "2008 TC4"
    assert param_text("q", " x ", 50) == " x "   # PHP's trim() keeps non-breaking spaces


def test_user_input_never_becomes_sql_text():
    aq = approach_query("2026-10-01", None, "true", "50%_off\\' OR 1=1", "distance", "desc")
    assert aq.where_sql == ("WHERE ca.approach_date >= :start AND a.is_potentially_hazardous = :hazardous "
                            "AND a.name LIKE :q ESCAPE '\\\\'")
    assert aq.params == {"start": "2026-10-01", "hazardous": 1, "q": "%50\\%\\_off\\\\' OR 1=1%"}
    assert aq.order_sql == "ca.miss_distance_km DESC, ca.id ASC"
    assert like_pattern("a_b") == "%a\\_b%"


def test_size_histogram_bins():
    sizes = [{"est_diameter_max_m": d, "is_potentially_hazardous": h}
             for d, h in [(5.0, 0), (10.0, 0), (24.9, 0), (140.0, 1), (1000.0, 1), (None, 0)]]
    bins = size_histogram(sizes)
    assert [b["label"] for b in bins] == ["0–10 m", "10–25 m", "25–50 m", "50–100 m", "100–250 m", "250–500 m",
                                         "500–1000 m", "≥ 1000 m"]
    assert [b["count"] for b in bins] == [1, 2, 0, 0, 1, 0, 0, 1]
    assert [b["hazardous"] for b in bins] == [0, 0, 0, 0, 1, 0, 0, 1]


def test_exports_cannot_run_spreadsheet_formulas():
    assert [safe_cell(v) for v in ["=1+1", "+cmd", "-2+3", "@SUM(A1)", "(2008 TC4)", True, -5.0]] == \
        ["'=1+1", "'+cmd", "'-2+3", "'@SUM(A1)", "(2008 TC4)", "yes", -5.0]
    rows = [{"name": "=HYPERLINK(\"http://x\")", "miss_distance_km": 1.5}]
    text = to_csv(rows, ["name", "miss_distance_km"]).decode("utf-8-sig")
    assert list(csv.reader(io.StringIO(text))) == [["name", "miss_distance_km"], ["'=HYPERLINK(\"http://x\")", "1.5"]]
    ws = load_workbook(io.BytesIO(to_xlsx(rows, ["name", "miss_distance_km"], "t"))).active
    assert ws["A2"].data_type == "s" and ws["A2"].value.startswith("'=")
    assert ws["B2"].value == 1.5


def test_dashboard_page_and_assets_are_served_but_nothing_else():
    page = client.get("/")
    assert page.status_code == 200 and "<title>Asteroid Watch</title>" in page.text
    assert page.headers["Cache-Control"] == "no-cache"
    assert client.get("/assets/js/main.js").headers["Cache-Control"] == "no-cache"
    for path in ["/assets/../../.env", "/assets/%2e%2e/%2e%2e/.env", "/.env", "/dashboard/index.html",
                 "/legacy/php/includes/db.php", "/pipeline/config.py"]:
        assert client.get(path).status_code == 404, path


def test_docs_load_only_local_files():
    docs = client.get("/docs")
    assert docs.status_code == 200
    urls = set(re.findall(r'(?:src|href)="([^"]+)"', docs.text))
    assert urls == {"/static/swagger-ui/swagger-ui-bundle.js", "/static/swagger-ui/swagger-ui.css",
                    "/assets/favicon.svg"}
    for url in urls:
        assert client.get(url).status_code == 200, url


def test_vendored_swagger_matches_recorded_hashes():
    readme = (SWAGGER / "README.md").read_text()
    rows = re.findall(r"\| `([^`]+)` \| `[^`]+` \| [\d,]+ \| `([0-9a-f]{64})` \|", readme)
    assert {name for name, _ in rows} == {"swagger-ui-bundle.js", "swagger-ui.css", "LICENSE", "NOTICE"}
    for name, sha in rows:
        assert hashlib.sha256((SWAGGER / name).read_bytes()).hexdigest() == sha, name


def test_openapi_lists_every_endpoint():
    paths = set(json.loads(client.get("/openapi.json").text)["paths"])
    assert paths == {"/api/summary", "/api/approaches", "/api/asteroid", "/api/trends", "/api/runs",
                     "/api/approaches.csv", "/api/approaches.xlsx", "/api/runs.csv", "/api/runs.xlsx"}
